from django.shortcuts import render
from ..forms import *
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.utils import timezone
from django.db.models import IntegerField
from django.db.models.functions import Cast

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment
from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd
from django.http import HttpResponse
from datetime import datetime

def cadastroAulasProcessoDigial(request):
    if request.method == 'POST':
        form = FormCadastroAulasProcessoDigital(request.POST)
        if form.is_valid():
            form.save()
            context = {
                'titulo': 'Treinamento para Utilização do Processo Digital',    
                'subtitulo': 'Subsecretaria de Tecnologia da Informação e Comunicação',
                'mensagem': "<span class='text-success'><i class='fa-solid fa-circle-check me-2'></i>Formulário enviado com sucesso!</span>"
            }
            return render(request, 'formfacil/formfacil_success.html', context)
    else:
        form = FormCadastroAulasProcessoDigital()
        
    context = {
            'form': form,
            'titulo': 'Treinamento para Utilização do Processo Digital',    
            'subtitulo': 'Subsecretaria de Tecnologia da Informação e Comunicação',
            'mensagem': (
                "No período de 09 a 13 de setembro, estaremos realizando o treinamento para Utilização do Processo Digital, que será implantado no dia 08 de outubro em toda a prefeitura. Para se inscrever, clique no melhor dia e horário para participar, preencha os dados solicitados e garanta a sua participação.<br><br>"
                "Local: Av. Alberto Braune, 223 - 2o. andar - Centro. (Auditório da Secretaria de Ciência e Tecnologia)"

            )
        }
    return render(request, 'formfacil/formfacil_form.html', context)

def visualizarDados_Aulas_Processo_Digital(request):
    hoje = timezone.now().date()
    # print(f"Data atual: {hoje.strftime('%d')}")
    hoje_int = int(hoje.strftime('%d'))

    cadastros = Opcao_Turmas.objects.annotate(
        valor_inteiro=Cast('dia_do_mes', IntegerField())
    ).filter(valor_inteiro__gte=hoje_int-1)

    # Preparar o contexto com os registros filtrados
    context = {
        'cadastros': cadastros,
    }
    
    return render(request, 'formfacil/visualizar_dados.html', context)

def exportar_aulas_processo_digital_to_excel(request):
    # Create a new workbook and select the active worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = 'Treinamento para Utilização do Processo Digital'

    # Load data from the database
    cadastros = Cadastro_Aulas_Processo_Digital.objects.all()
    ws.append(['Nome', 'Matrícula', 'Secretaria', 'Setor', 'Telefone', 'Turma escolhida', 'Data de inclusão'])
    # Write the data to the worksheet
    for item in cadastros:
        ws.append([
            item.nome,
            item.matricula,
            item.secretaria,
            item.setor,            
            item.telefone,
            item.turma_escolhida.__str__(),
            item.dt_registro.astimezone().strftime('%d/%m/%Y %H:%M:%S')
        ])
    # for r in dataframe_to_rows(df, index=False, header=True):
        # ws.append(r)

    # Create a table
    table = Table(displayName=f"TreinamentoProcessoDigital", ref=ws.dimensions)
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=True)
    table.tableStyleInfo = style
    ws.add_table(table)

    # Save the workbook to a BytesIO object
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    # Set up the response
    response = HttpResponse(content_type='application/vnd.ms-excel')
    response['Content-Disposition'] = 'attachment; filename=AulasProcessoDigital.xlsx'
    response.write(output.getvalue())
    return response

def logCadastrosRepetidos(request):
    # Data específica para filtrar os registros
    data_especifica = timezone.datetime(2024, 9, 16)
    
    # Obtém as matrículas que têm registros repetidos a partir da data especificada
    matriculas_repetidas = (
        Cadastro_Aulas_Processo_Digital.objects
        .filter(dt_registro__gte=data_especifica)  # Filtra registros a partir da data especificada
        .filter(~Q(matricula=''))  # Filtra apenas os que têm matrícula
        .values('matricula')  # Agrupa por matrícula
        .annotate(total=Count('matricula'))  # Conta quantas vezes cada matrícula aparece
        .filter(total__gt=1)  # Filtra apenas as matrículas que aparecem mais de uma vez
        .values_list('matricula', flat=True)  # Obtém uma lista de matrículas
    )

    # Obtém os nomes que têm registros repetidos a partir da data especificada (quando não têm matrícula)
    nomes_repetidos = (
        Cadastro_Aulas_Processo_Digital.objects
        .filter(dt_registro__gte=data_especifica)  # Filtra registros a partir da data especificada
        .filter(Q(matricula=''))  # Filtra apenas os que não têm matrícula
        .values('nome')  # Agrupa por nome
        .annotate(total=Count('nome'))  # Conta quantas vezes cada nome aparece
        .filter(total__gt=1)  # Filtra apenas os nomes que aparecem mais de uma vez
        .values_list('nome', flat=True)  # Obtém uma lista de nomes
    )

    # Obtém todos os registros para as matrículas repetidas ou nomes repetidos a partir da data especificada
    cadastros_repetidos = Cadastro_Aulas_Processo_Digital.objects.filter(
        Q(matricula__in=matriculas_repetidas) | Q(nome__in=nomes_repetidos),
        dt_registro__gte=data_especifica  # Filtra registros a partir da data especificada
    )

    # Renderiza o template com os cadastros repetidos
    return render(request, 'formfacil/cadastrosrepetidos.html', {'cadastros': cadastros_repetidos, 'total': cadastros_repetidos.count()})



def inscricaoDecretosPortariaAtosPrefeito(request):
    if request.method == 'POST':
        form = InscricaoDecretosPortariaAtosPrefeitoForm(request.POST)
        if form.is_valid():
            print("Form is valid")
            form.save()
            print("Form saved")
            context = {
                'titulo': 'Inscrição Decretos e Portaria Atos do Prefeito',    
                'subtitulo': 'Subsecretaria de Tecnologia da Informação e Comunicação',
                'mensagem': "<span class='text-success'><i class='fa-solid fa-circle-check me-2'></i>Formulário enviado com sucesso!</span>"
            }
            return render(request, 'formfacil/formfacil_success.html', context)
    else:
        form = InscricaoDecretosPortariaAtosPrefeitoForm()
        
    context = {
            'form': form,
            'titulo': 'Inscrição Decretos e Portaria Atos do Prefeito',    
            'subtitulo': 'Subsecretaria de Tecnologia da Informação e Comunicação',
            'mensagem': ''
    }
    return render(request, 'formfacil/formfacil_form.html', context)